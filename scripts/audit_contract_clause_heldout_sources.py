from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from docx import Document


ROOT = Path(
    "data/evaluation/contract_review_clause_retrieval/heldout/contracts"
)


EXPECTED_FILES = {
    "H01_Backend_Developer.docx",
    "H02_UI_UX_Designer.docx",
    "H03_Network_Engineer.docx",
    "H04_Customer_Support.docx",
    "H05_Warehouse_Supervisor.docx",
    "H06_Security_Analyst.docx",
    "H07_HR_Specialist.docx",
    "H08_Sales_Executive.docx",
    "H09_DevOps_Engineer.docx",
    "H10_Office_Administrator.docx",
    "H11_QA_Engineer.docx",
    "H12_Finance_Assistant.docx",
    "H13_Call_Center_Agent.docx",
    "H14_Procurement_Specialist.docx",
    "H15_Data_Engineer.docx",
    "H16_Operations_Coordinator.docx",
}


# These checks are only guards against obvious source-generation mistakes.
# They are NOT the gold annotation logic.
FORBIDDEN_BY_MISSING_CATEGORY = {
    "H10_Office_Administrator.docx": {
        "category": "probation",
        "phrases": [
            "thử việc",
            "lương thử việc",
            "giai đoạn đánh giá ban đầu",
            "giai đoạn làm quen công việc",
            "giai đoạn đầu",
        ],
    },
    "H12_Finance_Assistant.docx": {
        "category": "termination",
        "phrases": [
            "chấm dứt",
            "nghỉ việc",
            "báo trước",
            "kết thúc quan hệ lao động",
            "kết thúc hợp đồng",
        ],
    },
    "H15_Data_Engineer.docx": {
        "category": "salary",
        "phrases": [
            "mức lương",
            "tiền lương",
            "lương thử việc",
            "phụ cấp",
            "ngày trả lương",
            "thu nhập cố định",
            "khoản thanh toán cố định",
        ],
    },
    "H16_Operations_Coordinator.docx": {
        "category": "working_time",
        "phrases": [
            "thời giờ làm việc",
            "giờ/ngày",
            "ngày/tuần",
            "lịch làm việc",
            "phân ca",
            "nghỉ giữa ca",
            "nghỉ hằng tuần",
            "làm thêm",
            "on-call",
        ],
    },
}


EXPECTED_TABLES = {
    "H03_Network_Engineer.docx",
    "H05_Warehouse_Supervisor.docx",
    "H08_Sales_Executive.docx",
    "H11_QA_Engineer.docx",
    "H13_Call_Center_Agent.docx",
    "H16_Operations_Coordinator.docx",
}


def normalize(text: str) -> str:
    text = text.lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(
        ch for ch in text
        if unicodedata.category(ch) != "Mn"
    )
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_text(doc: Document) -> str:
    parts: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]
            parts.append(" | ".join(cells))

    return "\n".join(parts)


def numbered_headings(text: str) -> list[str]:
    result = []

    for line in text.splitlines():
        if re.match(
            r"^\s*Điều\s+\d+[.:]",
            line,
            flags=re.IGNORECASE,
        ):
            result.append(line.strip())

    return result


def main() -> None:
    errors: list[str] = []
    warnings: list[str] = []

    files = sorted(ROOT.glob("*.docx"))
    names = {p.name for p in files}

    print("Held-out source audit")
    print("=====================")
    print(f"DOCX count: {len(files)}")

    missing_files = EXPECTED_FILES - names
    extra_files = names - EXPECTED_FILES

    if missing_files:
        errors.append(
            "Missing files: "
            + ", ".join(sorted(missing_files))
        )

    if extra_files:
        errors.append(
            "Unexpected files: "
            + ", ".join(sorted(extra_files))
        )

    if len(files) != 16:
        errors.append(
            f"Expected exactly 16 DOCX files, got {len(files)}"
        )

    print()

    for path in files:
        doc = Document(path)
        text = extract_text(doc)
        headings = numbered_headings(text)

        print(
            f"{path.name}: "
            f"headings={len(headings)}, "
            f"tables={len(doc.tables)}, "
            f"chars={len(text)}"
        )

        if len(headings) < 7 or len(headings) > 8:
            warnings.append(
                f"{path.name}: expected roughly 7-8 numbered "
                f"sections from frozen specs, got {len(headings)}"
            )

        if path.name in EXPECTED_TABLES and not doc.tables:
            errors.append(
                f"{path.name}: expected at least one DOCX table"
            )

        rule = FORBIDDEN_BY_MISSING_CATEGORY.get(
            path.name
        )

        if rule:
            norm = normalize(text)

            for phrase in rule["phrases"]:
                if normalize(phrase) in norm:
                    errors.append(
                        f"{path.name}: missing-category "
                        f"{rule['category']} contains forbidden "
                        f"phrase {phrase!r}"
                    )

    print()
    print("Missing-category guards")
    print("-----------------------")

    for filename, rule in (
        FORBIDDEN_BY_MISSING_CATEGORY.items()
    ):
        print(
            f"{filename}: "
            f"{rule['category']} -> checked"
        )

    if warnings:
        print()
        print("WARNINGS")
        print("--------")
        for item in warnings:
            print(f"[WARN] {item}")

    print()

    if errors:
        print("FAILED")
        print("------")
        for item in errors:
            print(f"[ERROR] {item}")

        raise SystemExit(1)

    print("PASS")
    print("----")
    print(
        "Basic held-out source audit passed. "
        "Manual semantic review is still required."
    )


if __name__ == "__main__":
    main()
