#!/usr/bin/env python3
"""Real HTTP E2E for contract review using a persistent application database."""

from __future__ import annotations

import argparse
import json
import os
import re
import signal
import socket
import subprocess
import sys
import tempfile
import time
import unicodedata
import zipfile
from pathlib import Path

import httpx
import pymupdf
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "samples/contract-review/sample_labor_contract.docx"
PDF_SAMPLE = ROOT / "samples/contract-review/sample_labor_contract.pdf"


def free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def wait_live(base_url: str, timeout: float = 30) -> None:
    deadline = time.time() + timeout
    with httpx.Client(trust_env=False) as client:
        while time.time() < deadline:
            try:
                if client.get(f"{base_url}/live", timeout=1).status_code == 200:
                    return
            except httpx.HTTPError:
                pass
            time.sleep(0.2)
    raise RuntimeError("Backend runtime did not become ready.")


def start_server(
    port: int,
    database_path: Path,
    log_path: Path,
    *,
    require_authority_approval: bool,
) -> subprocess.Popen[str]:
    env = os.environ.copy()
    env.update({
        "PYTHONPATH": str(ROOT / "backend"),
        "DATABASE_URL": f"sqlite+pysqlite:///{database_path}",
        "DATABASE_AUTO_CREATE": "true",
        "PASSWORD_PBKDF2_ITERATIONS": "100000",
        "LEGAL_CHUNKS_PATH": str(ROOT / "data/releases/labor-law-canonical-word-20260804-164432-candidate/canonical_chunks.jsonl"),
        "CORPUS_REQUIRE_AUTHORITY_APPROVAL": (
            "true" if require_authority_approval else "false"
        ),
    })
    log = log_path.open("a", encoding="utf-8")
    process = subprocess.Popen(
        [
            sys.executable,
            "-m",
            "uvicorn",
            "app.main:app",
            "--host",
            "127.0.0.1",
            "--port",
            str(port),
            "--log-level",
            "warning",
        ],
        cwd=ROOT,
        env=env,
        stdout=log,
        stderr=subprocess.STDOUT,
        text=True,
    )
    return process


def stop_server(process: subprocess.Popen[str]) -> None:
    if process.poll() is not None:
        return
    process.send_signal(signal.SIGTERM)
    try:
        process.wait(timeout=10)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=5)


def register(client: httpx.Client, email: str) -> str:
    response = client.post(
        "/auth/register",
        json={
            "email": email,
            "password": "MatKhauAnToan123",
            "display_name": email.split("@", 1)[0],
        },
    )
    assert response.status_code == 201, response.text
    csrf = client.cookies.get("legal_rag_csrf")
    assert csrf, "missing csrf cookie"
    return csrf


def assert_markers(review: dict[str, object]) -> None:
    findings = review["findings"]
    assert isinstance(findings, list) and len(findings) == 4
    assert {item["category"] for item in findings} == {
        "probation", "salary", "working_time", "termination"
    }
    for finding in findings:
        sources = finding["sources"]
        assert sources, finding
        source_ids = {source["source_id"] for source in sources}
        article_codes = [source["article_code"] for source in sources]
        markers = set(re.findall(r"\[(S\d+)\]", finding["analysis"]))
        assert markers, finding["analysis"]
        assert markers == source_ids, (markers, source_ids)
        assert len(article_codes) == len(set(article_codes)), article_codes
        assert finding["contract_excerpt"]
        assert finding["recommendation"]

    by_category = {item["category"]: item for item in findings}
    assert by_category["probation"]["severity"] == "attention"
    probation_excerpt = "".join(
        character
        for character in unicodedata.normalize(
            "NFD", by_category["probation"]["contract_excerpt"].casefold()
        ).replace("đ", "d")
        if unicodedata.category(character) != "Mn"
    )
    salary_excerpt = "".join(
        character
        for character in unicodedata.normalize(
            "NFD", by_category["salary"]["contract_excerpt"].casefold()
        ).replace("đ", "d")
        if unicodedata.category(character) != "Mn"
    )
    termination_excerpt = "".join(
        character
        for character in unicodedata.normalize(
            "NFD", by_category["termination"]["contract_excerpt"].casefold()
        ).replace("đ", "d")
        if unicodedata.category(character) != "Mn"
    )
    assert "75 ngay" in probation_excerpt
    assert by_category["salary"]["severity"] == "info"
    assert "12.000.000 dong" in salary_excerpt
    assert by_category["working_time"]["severity"] == "attention"
    assert "45 giờ/tuần" in by_category["working_time"]["analysis"]
    assert by_category["termination"]["severity"] == "warning"
    assert "10 ngay" in termination_excerpt


def create_adversarial_docx(path: Path) -> None:
    document = Document()
    for paragraph in (
        "HỢP ĐỒNG LAO ĐỘNG XÁC ĐỊNH THỜI HẠN 24 tháng",
        "Thử việc 30 ngày; khoản thanh toán hoàn tất trong 90 ngày.",
        "Tiền lương 12.000.000 đồng, trả vào ngày 05 hằng tháng.",
        "Thời giờ làm việc 8 giờ/ngày, 48 giờ/tuần.",
        "Khi chấm dứt hợp đồng phải báo trước 45 ngày; nghỉ phép riêng 1 ngày.",
    ):
        document.add_paragraph(paragraph)
    document.save(path)


def create_unrelated_numbers_docx(path: Path) -> None:
    document = Document()
    for paragraph in (
        "Số hợp đồng 123/2026.",
        "Người lao động: Nguyễn Văn A.",
        "Căn cước công dân số 012345678901, cấp ngày 01/01/2026.",
    ):
        document.add_paragraph(paragraph)
    document.save(path)


def create_blank_pdf(path: Path, page_count: int) -> None:
    document = pymupdf.open()
    try:
        for _ in range(page_count):
            document.new_page()
        document.save(path)
    finally:
        document.close()


def create_docx_zip_bomb_probe(path: Path) -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/large.xml", b"0" * (50 * 1024 * 1024 + 1))


def upload_contract(
    client: httpx.Client,
    csrf: str,
    path: Path,
    content_type: str,
) -> httpx.Response:
    with path.open("rb") as source:
        return client.post(
            "/contract-reviews",
            headers={"X-CSRF-Token": csrf},
            data={"method": "sparse"},
            files={"file": (path.name, source, content_type)},
        )


def upload_bytes(
    client: httpx.Client,
    csrf: str | None,
    *,
    filename: str,
    content_type: str,
    data: bytes,
) -> httpx.Response:
    headers = {"X-CSRF-Token": csrf} if csrf else {}
    return client.post(
        "/contract-reviews",
        headers=headers,
        data={"method": "sparse"},
        files={"file": (filename, data, content_type)},
    )


def run() -> dict[str, object]:
    with tempfile.TemporaryDirectory(prefix="contract-review-e2e-") as temp_name:
        temp = Path(temp_name)
        database_path = temp / "runtime.db"
        server_log = temp / "server.log"
        adversarial_path = temp / "adversarial_numbers_contract.docx"
        unrelated_path = temp / "unrelated_numbers_contract.docx"
        scan_pdf_path = temp / "scan.pdf"
        too_many_pages_path = temp / "too-many-pages.pdf"
        zip_bomb_path = temp / "expanded-too-large.docx"
        create_adversarial_docx(adversarial_path)
        create_unrelated_numbers_docx(unrelated_path)
        create_blank_pdf(scan_pdf_path, 1)
        create_blank_pdf(too_many_pages_path, 251)
        create_docx_zip_bomb_probe(zip_bomb_path)
        port = free_port()
        base_url = f"http://127.0.0.1:{port}/api"
        process = start_server(
            port,
            database_path,
            server_log,
            require_authority_approval=True,
        )
        steps: list[str] = []
        try:
            wait_live(base_url)
            steps.append("PASS 1: backend HTTP runtime ready")
            user_a = httpx.Client(base_url=base_url, timeout=120, trust_env=False)
            user_b = httpx.Client(base_url=base_url, timeout=120, trust_env=False)
            csrf_a = register(user_a, "runtime-contract-a@example.test")
            csrf_b = register(user_b, "runtime-contract-b@example.test")
            steps.append("PASS 2: registered two users with real session cookies")

            blocked = upload_contract(
                user_a,
                csrf_a,
                SAMPLE,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            assert blocked.status_code == 503, blocked.text
            assert "authority_review_pending" in blocked.text
            assert user_a.get("/contract-reviews").json()["total"] == 0
            steps.append("PASS 3: authority gate blocked a new legal report before persistence")

            stop_server(process)
            process = start_server(
                port,
                database_path,
                server_log,
                require_authority_approval=False,
            )
            wait_live(base_url)

            anonymous = httpx.Client(base_url=base_url, timeout=120, trust_env=False)
            invalid_cases = [
                (
                    upload_bytes(
                        anonymous,
                        None,
                        filename="contract.txt",
                        content_type="text/plain",
                        data=b"plain text",
                    ),
                    401,
                    None,
                ),
                (
                    upload_bytes(
                        user_a,
                        "invalid-csrf",
                        filename="contract.txt",
                        content_type="text/plain",
                        data=b"plain text",
                    ),
                    403,
                    None,
                ),
                (
                    upload_bytes(
                        user_a,
                        csrf_a,
                        filename="contract.txt",
                        content_type="text/plain",
                        data=b"plain text",
                    ),
                    422,
                    "PDF hoặc DOCX",
                ),
                (
                    upload_bytes(
                        user_a,
                        csrf_a,
                        filename="empty.docx",
                        content_type="application/octet-stream",
                        data=b"",
                    ),
                    422,
                    "trống",
                ),
                (
                    upload_bytes(
                        user_a,
                        csrf_a,
                        filename="corrupt.docx",
                        content_type="application/octet-stream",
                        data=b"not-a-docx",
                    ),
                    422,
                    "bị hỏng",
                ),
                (
                    upload_bytes(
                        user_a,
                        csrf_a,
                        filename="fake.pdf",
                        content_type="application/pdf",
                        data=b"%PDF-1.7 fake",
                    ),
                    422,
                    "bị hỏng",
                ),
                (
                    upload_bytes(
                        user_a,
                        csrf_a,
                        filename="scan.pdf",
                        content_type="application/pdf",
                        data=scan_pdf_path.read_bytes(),
                    ),
                    422,
                    "chưa hỗ trợ OCR",
                ),
                (
                    upload_bytes(
                        user_a,
                        csrf_a,
                        filename="too-many-pages.pdf",
                        content_type="application/pdf",
                        data=too_many_pages_path.read_bytes(),
                    ),
                    422,
                    "250 trang",
                ),
                (
                    upload_bytes(
                        user_a,
                        csrf_a,
                        filename="expanded-too-large.docx",
                        content_type="application/octet-stream",
                        data=zip_bomb_path.read_bytes(),
                    ),
                    422,
                    "giải nén vượt giới hạn",
                ),
                (
                    upload_bytes(
                        user_a,
                        csrf_a,
                        filename="oversized.docx",
                        content_type="application/octet-stream",
                        data=b"x" * (10 * 1024 * 1024 + 1),
                    ),
                    422,
                    "10 MB",
                ),
                (
                    upload_bytes(
                        user_a,
                        csrf_a,
                        filename=PDF_SAMPLE.name,
                        content_type="image/png",
                        data=PDF_SAMPLE.read_bytes(),
                    ),
                    422,
                    "MIME type",
                ),
            ]
            for invalid_response, expected_status, expected_text in invalid_cases:
                assert invalid_response.status_code == expected_status, invalid_response.text
                if expected_text:
                    assert expected_text in invalid_response.text, invalid_response.text
            anonymous.close()
            assert user_a.get("/contract-reviews").json()["total"] == 0
            steps.append("PASS 4: rejected 11 invalid, unsafe, or unauthorized uploads")

            response = upload_contract(
                user_a,
                csrf_a,
                SAMPLE,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            assert response.status_code == 201, response.text
            review = response.json()
            review_id = review["id"]
            assert review["extracted_character_count"] > 500
            assert_markers(review)
            assert "Có 2 nhóm cần kiểm tra" in review["summary"]
            assert "1 nhóm cần ưu tiên kiểm tra" in review["summary"]
            steps.append("PASS 5: uploaded real DOCX and created four grounded findings")

            adversarial_response = upload_contract(
                user_a,
                csrf_a,
                adversarial_path,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            assert adversarial_response.status_code == 201, adversarial_response.text
            adversarial_review = adversarial_response.json()
            adversarial_id = adversarial_review["id"]
            adversarial_findings = {
                item["category"]: item for item in adversarial_review["findings"]
            }
            assert "30 ngày" in adversarial_findings["probation"]["analysis"]
            assert "90 ngày" not in adversarial_findings["probation"]["analysis"]
            assert "8 giờ/ngày" in adversarial_findings["working_time"]["analysis"]
            assert "48 giờ/ngày" not in adversarial_findings["working_time"]["analysis"]
            assert "45 ngày" in adversarial_findings["termination"]["analysis"]
            assert "1 ngày" not in adversarial_findings["termination"]["analysis"]
            steps.append("PASS 6: adversarial DOCX kept each number bound to its legal term")

            unrelated_response = upload_contract(
                user_a,
                csrf_a,
                unrelated_path,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            assert unrelated_response.status_code == 201, unrelated_response.text
            unrelated_review = unrelated_response.json()
            unrelated_id = unrelated_review["id"]
            assert "4 nhóm chưa tìm thấy" in unrelated_review["summary"]
            assert all(
                item["contract_excerpt"].startswith("Chưa tìm thấy")
                for item in unrelated_review["findings"]
            )
            steps.append("PASS 7: unrelated numbers did not invent contract clauses")

            pdf_response = upload_contract(
                user_a,
                csrf_a,
                PDF_SAMPLE,
                "application/pdf",
            )
            assert pdf_response.status_code == 201, pdf_response.text
            pdf_review = pdf_response.json()
            pdf_id = pdf_review["id"]
            assert pdf_review["extracted_character_count"] > 300
            assert_markers(pdf_review)
            steps.append("PASS 8: uploaded a real text-layer PDF through the HTTP API")

            listing = user_a.get("/contract-reviews")
            assert listing.status_code == 200 and listing.json()["total"] == 4
            sample_item = next(
                item for item in listing.json()["items"] if item["id"] == review_id
            )
            assert sample_item["attention_count"] == 2
            assert sample_item["warning_count"] == 1
            detail = user_a.get(f"/contract-reviews/{review_id}")
            assert detail.status_code == 200
            assert detail.json()["file_sha256"] == review["file_sha256"]
            steps.append("PASS 9: list and detail APIs returned all persisted reports")

            assert user_b.get(f"/contract-reviews/{review_id}").status_code == 404
            assert user_b.delete(
                f"/contract-reviews/{review_id}", headers={"X-CSRF-Token": csrf_b}
            ).status_code == 404
            steps.append("PASS 10: cross-user read and delete were denied with 404")

            stop_server(process)
            process = start_server(
                port,
                database_path,
                server_log,
                require_authority_approval=False,
            )
            wait_live(base_url)
            persisted = user_a.get(f"/contract-reviews/{review_id}")
            assert persisted.status_code == 200, persisted.text
            assert_markers(persisted.json())
            assert user_a.get(f"/contract-reviews/{adversarial_id}").status_code == 200
            assert user_a.get(f"/contract-reviews/{unrelated_id}").status_code == 200
            assert user_a.get(f"/contract-reviews/{pdf_id}").status_code == 200
            assert user_a.get("/auth/me").status_code == 200
            steps.append("PASS 11: session and all reports survived backend process restart")

            for stored_id in (review_id, adversarial_id, unrelated_id, pdf_id):
                deleted = user_a.delete(
                    f"/contract-reviews/{stored_id}",
                    headers={"X-CSRF-Token": csrf_a},
                )
                assert deleted.status_code == 204, deleted.text
                assert user_a.get(f"/contract-reviews/{stored_id}").status_code == 404
            assert user_a.get("/contract-reviews").json()["total"] == 0
            steps.append("PASS 12: owner deleted all reports and no data remained")

            user_a.close()
            user_b.close()
            return {
                "status": "PASS",
                "runtime": "full FastAPI application over real HTTP",
                "database": "persistent SQLite file",
                "sample": str(SAMPLE.relative_to(ROOT)),
                "steps": steps,
                "review_id": review_id,
                "finding_summary": [
                    {
                        "category": item["category"],
                        "severity": item["severity"],
                        "article_codes": [
                            source["article_code"] for source in item["sources"]
                        ],
                        "analysis": item["analysis"],
                    }
                    for item in review["findings"]
                ],
                "server_log_tail": server_log.read_text(encoding="utf-8", errors="replace").splitlines()[-20:],
            }
        finally:
            stop_server(process)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = run()
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
