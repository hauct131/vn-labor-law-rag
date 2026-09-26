from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(
    "data/evaluation/contract_review_clause_retrieval/heldout/contracts"
)
ROOT.mkdir(parents=True, exist_ok=True)


def setup_document() -> Document:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    return doc


def title(doc: Document, value: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r = p.add_run("HỢP ĐỒNG LAO ĐỘNG")
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(value)
    r.bold = True


def preamble(doc: Document, role: str) -> None:
    doc.add_paragraph(
        "BÊN A: CÔNG TY TNHH DEMO HELDOUT VIỆT NAM"
    )
    doc.add_paragraph(
        "BÊN B: NGƯỜI LAO ĐỘNG"
    )
    doc.add_paragraph(
        f"Hai bên thống nhất ký hợp đồng lao động cho vị trí {role} "
        "với các nội dung sau đây."
    )


def heading(doc: Document, number: int, name: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"Điều {number}. {name}")
    r.bold = True


def paragraphs(
    doc: Document,
    *items: str,
) -> None:
    for item in items:
        doc.add_paragraph(item)


def table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    t = doc.add_table(
        rows=1,
        cols=len(headers),
    )
    t.style = "Table Grid"

    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = value

    for row in rows:
        cells = t.add_row().cells

        for i, value in enumerate(row):
            cells[i].text = value


def save(doc: Document, filename: str) -> None:
    path = ROOT / filename
    doc.save(path)
    print(path)


# ============================================================
# H01
# ============================================================

def h01() -> None:
    doc = setup_document()
    title(doc, "H01 - Backend Developer")
    preamble(doc, "Backend Developer")

    heading(doc, 1, "Công việc và địa điểm làm việc")
    paragraphs(
        doc,
        "Người lao động phát triển API, bảo trì các dịch vụ backend "
        "và phối hợp với nhóm frontend trong quá trình tích hợp.",
        "Địa điểm làm việc chính tại văn phòng công ty tại Thành phố Hồ Chí Minh.",
    )

    heading(doc, 2, "Thử việc")
    paragraphs(
        doc,
        "Người lao động thực hiện thử việc trong thời gian 60 ngày.",
        "Trong thời gian thử việc, người lao động thực hiện các nhiệm vụ "
        "của vị trí Backend Developer.",
        "Kết quả thử việc được đánh giá vào cuối thời gian thử việc.",
    )

    heading(doc, 3, "Thời hạn hợp đồng")
    paragraphs(
        doc,
        "Hợp đồng xác định thời hạn 24 tháng kể từ ngày 01/11/2026 "
        "đến hết ngày 31/10/2028.",
    )

    heading(doc, 4, "Tiền lương")
    paragraphs(
        doc,
        "Mức lương theo công việc là 22.000.000 đồng/tháng.",
        "Phụ cấp ăn trưa là 700.000 đồng/tháng.",
        "Tiền lương được trả vào ngày 05 hằng tháng bằng chuyển khoản ngân hàng.",
    )

    heading(doc, 5, "Thời giờ làm việc và nghỉ ngơi")
    paragraphs(
        doc,
        "Thời giờ làm việc là 08 giờ/ngày, 05 ngày/tuần, "
        "từ thứ Hai đến thứ Sáu.",
        "Thời gian làm việc từ 08:30 đến 17:30, nghỉ trưa một giờ.",
        "Người lao động được nghỉ hằng tuần vào thứ Bảy và Chủ nhật.",
    )

    heading(doc, 6, "Bảo hiểm và nghỉ phép")
    paragraphs(
        doc,
        "Người lao động tham gia BHXH, BHYT và BHTN theo quy định.",
        "Nghỉ phép năm thực hiện theo chính sách của công ty và pháp luật.",
    )

    heading(doc, 7, "Bảo mật")
    paragraphs(
        doc,
        "Người lao động có trách nhiệm bảo mật mã nguồn, dữ liệu "
        "và thông tin kỹ thuật của công ty.",
    )

    heading(doc, 8, "Chấm dứt hợp đồng")
    paragraphs(
        doc,
        "Việc chấm dứt hợp đồng được thực hiện theo căn cứ và trình tự "
        "đã thỏa thuận và quy định pháp luật.",
        "Bên thực hiện quyền đơn phương có trách nhiệm tuân thủ thời hạn báo trước.",
        "Người lao động phải bàn giao mã nguồn, tài khoản và tài liệu "
        "trước ngày làm việc cuối cùng.",
    )

    save(doc, "H01_Backend_Developer.docx")


# ============================================================
# H02
# ============================================================

def h02() -> None:
    doc = setup_document()
    title(doc, "H02 - UI UX Designer")
    preamble(doc, "UI/UX Designer")

    heading(doc, 1, "Phạm vi công việc")
    paragraphs(
        doc,
        "Người lao động thiết kế giao diện, prototype và duy trì design system "
        "cho các sản phẩm số của công ty.",
    )

    heading(doc, 2, "Giai đoạn ban đầu")
    paragraphs(
        doc,
        "Trong 60 ngày đầu kể từ ngày nhận việc, người lao động thực hiện "
        "đầy đủ nhiệm vụ của vị trí UI/UX Designer.",
        "Cuối giai đoạn này, công ty đánh giá khả năng đáp ứng yêu cầu "
        "về chất lượng thiết kế, phối hợp và tiến độ.",
        "Nếu đạt yêu cầu, người lao động tiếp tục thực hiện hợp đồng "
        "theo các điều khoản đã thống nhất.",
    )

    heading(doc, 3, "Thời hạn")
    paragraphs(
        doc,
        "Thỏa thuận lao động có thời hạn 18 tháng kể từ ngày bắt đầu công việc.",
    )

    heading(doc, 4, "Quyền lợi")
    paragraphs(
        doc,
        "Khoản thanh toán cố định cho công việc là 19.500.000 đồng mỗi tháng.",
        "Người lao động được hỗ trợ thiết bị 500.000 đồng mỗi tháng.",
        "Các khoản trên được chuyển vào tài khoản cá nhân vào ngày 05 hằng tháng.",
    )

    heading(doc, 5, "Tổ chức công việc")
    paragraphs(
        doc,
        "Người lao động thực hiện nhiệm vụ từ 09:00 đến 18:00 "
        "từ thứ Hai đến thứ Sáu.",
        "Khoảng thời gian 12:00 đến 13:00 được dành cho nghỉ trưa.",
    )

    heading(doc, 6, "Sở hữu trí tuệ")
    paragraphs(
        doc,
        "Các bản thiết kế và tài sản sáng tạo hình thành trong quá trình "
        "thực hiện nhiệm vụ thuộc phạm vi quản lý của công ty.",
    )

    heading(doc, 7, "Khi quan hệ lao động kết thúc")
    paragraphs(
        doc,
        "Một bên có thể kết thúc quan hệ lao động trước thời hạn khi có căn cứ "
        "phù hợp với thỏa thuận và quy định áp dụng.",
        "Bên thực hiện việc kết thúc phải gửi thông báo trong thời gian "
        "phù hợp trước ngày dự kiến kết thúc.",
        "Người lao động hoàn tất bàn giao file thiết kế, tài liệu "
        "và quyền truy cập trước ngày cuối cùng.",
    )

    save(doc, "H02_UI_UX_Designer.docx")


# ============================================================
# H03
# ============================================================

def h03() -> None:
    doc = setup_document()
    title(doc, "H03 - Network Engineer")
    preamble(doc, "Network Engineer")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động giám sát mạng, cấu hình thiết bị và xử lý sự cố hạ tầng.",
    )

    heading(doc, 2, "Thử việc")
    paragraphs(
        doc,
        "Thời gian thử việc là 60 ngày.",
        "Kết quả được đánh giá căn cứ khả năng xử lý sự cố "
        "và vận hành hệ thống mạng.",
    )

    heading(doc, 3, "Tiền lương")
    paragraphs(
        doc,
        "Mức lương theo công việc là 21.000.000 đồng/tháng.",
        "Phụ cấp trực hệ thống là 1.200.000 đồng/tháng.",
        "Tiền lương được thanh toán vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Thời giờ làm việc")
    paragraphs(
        doc,
        "Thời giờ làm việc bình thường là 08 giờ/ngày và 05 ngày/tuần.",
        "Khung giờ hành chính là 08:30 đến 17:30.",
    )

    heading(doc, 5, "Lịch trực hệ thống")
    table(
        doc,
        ["Ca", "Khung giờ", "Ghi chú"],
        [
            ["A", "06:00-14:00", "Trực sự cố"],
            ["B", "14:00-22:00", "Trực sự cố"],
            ["C", "22:00-06:00", "Trực đêm"],
        ],
    )
    paragraphs(
        doc,
        "Lịch trực được phân công theo tuần.",
        "Người lao động phải có tối thiểu 08 giờ nghỉ trước khi "
        "được bố trí sang ca tiếp theo.",
    )

    heading(doc, 6, "Nghỉ phép và bảo hiểm")
    paragraphs(
        doc,
        "Các chế độ nghỉ phép năm và bảo hiểm thực hiện theo quy định.",
    )

    heading(doc, 7, "Thiết bị")
    paragraphs(
        doc,
        "Thiết bị mạng được giao phải được sử dụng đúng mục đích "
        "và bảo quản theo quy trình nội bộ.",
    )

    heading(doc, 8, "Chấm dứt hợp đồng")
    paragraphs(
        doc,
        "Việc chấm dứt hợp đồng được thực hiện theo căn cứ pháp luật "
        "và thỏa thuận của hai bên.",
        "Thời hạn báo trước áp dụng theo từng trường hợp.",
        "Người lao động phải bàn giao thiết bị và tài liệu kỹ thuật.",
    )

    save(doc, "H03_Network_Engineer.docx")


# ============================================================
# H04
# ============================================================

def h04() -> None:
    doc = setup_document()
    title(doc, "H04 - Customer Support")
    preamble(doc, "Customer Support")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động hỗ trợ khách hàng qua điện thoại, email và kênh chat.",
    )

    heading(doc, 2, "Giai đoạn làm quen công việc")
    paragraphs(
        doc,
        "Trong 30 ngày đầu kể từ ngày nhận việc, người lao động được hướng dẫn "
        "quy trình và thực hiện các yêu cầu hỗ trợ thực tế.",
        "Khả năng xử lý yêu cầu và chất lượng giao tiếp được đánh giá cuối giai đoạn.",
        "Kết quả đánh giá được sử dụng để quyết định việc tiếp tục bố trí "
        "người lao động vào vị trí.",
    )

    heading(doc, 3, "Thu nhập")
    paragraphs(
        doc,
        "Khoản thu nhập cố định là 13.500.000 đồng/tháng.",
        "Hỗ trợ điện thoại là 400.000 đồng/tháng.",
        "Khoản thu nhập được chuyển vào tài khoản vào ngày 07 hằng tháng.",
    )

    heading(doc, 4, "Lịch phục vụ khách hàng")
    paragraphs(
        doc,
        "Ca sáng từ 07:00 đến 15:00 và ca chiều từ 15:00 đến 23:00.",
        "Nhân sự được bố trí luân phiên giữa các ca theo lịch hằng tuần.",
        "Mỗi ca có thời gian nghỉ giữa ca theo lịch vận hành.",
    )

    heading(doc, 5, "Thiết bị")
    paragraphs(
        doc,
        "Tai nghe và máy tính xách tay do công ty cấp phải được hoàn trả "
        "khi hợp đồng chấm dứt.",
    )

    heading(doc, 6, "Bảo mật")
    paragraphs(
        doc,
        "Thông tin khách hàng và dữ liệu nội bộ phải được bảo mật.",
    )

    heading(doc, 7, "Kết thúc hợp đồng")
    paragraphs(
        doc,
        "Hợp đồng có thể kết thúc theo các căn cứ được pháp luật "
        "và thỏa thuận giữa hai bên cho phép.",
        "Bên thực hiện quyền kết thúc phải tuân thủ thời gian thông báo tương ứng.",
        "Người lao động phải hoàn tất việc bàn giao hồ sơ khách hàng.",
    )

    save(doc, "H04_Customer_Support.docx")


# ============================================================
# H05
# ============================================================

def h05() -> None:
    doc = setup_document()
    title(doc, "H05 - Warehouse Supervisor")
    preamble(doc, "Warehouse Supervisor")

    heading(doc, 1, "Trách nhiệm")
    paragraphs(
        doc,
        "Người lao động quản lý hoạt động kho, nhân sự và việc bàn giao hàng hóa.",
    )

    heading(doc, 2, "Thử việc")
    paragraphs(
        doc,
        "Người lao động thử việc 45 ngày và được đánh giá "
        "về năng lực điều phối kho.",
    )

    heading(doc, 3, "Chế độ")
    table(
        doc,
        ["Khoản", "Mức"],
        [
            ["Lương theo công việc", "17.000.000 đồng/tháng"],
            ["Phụ cấp trách nhiệm", "1.500.000 đồng/tháng"],
            ["Phụ cấp ca", "800.000 đồng/tháng"],
        ],
    )
    paragraphs(
        doc,
        "Các khoản được thanh toán vào ngày 07 hằng tháng bằng chuyển khoản.",
    )

    heading(doc, 4, "Tổ chức vận hành kho")
    paragraphs(
        doc,
        "Người lao động làm việc theo ca 12 giờ.",
        "Lịch ca được công bố theo tuần và người lao động làm việc "
        "04 ngày trong một tuần tiêu chuẩn.",
    )

    heading(doc, 5, "Nghỉ và chuyển ca")
    paragraphs(
        doc,
        "Người lao động được bố trí nghỉ giữa ca.",
        "Khoảng nghỉ tối thiểu giữa hai ca liên tiếp là 12 giờ.",
        "Ngày nghỉ hằng tuần được sắp xếp trong lịch phân ca.",
    )

    heading(doc, 6, "An toàn lao động")
    paragraphs(
        doc,
        "Người lao động phải tuân thủ quy trình an toàn và sử dụng "
        "trang thiết bị bảo hộ.",
    )

    heading(doc, 7, "Nghĩa vụ quản lý")
    paragraphs(
        doc,
        "Người lao động chịu trách nhiệm kiểm kê và báo cáo chênh lệch hàng hóa.",
    )

    heading(doc, 8, "Chấm dứt")
    paragraphs(
        doc,
        "Việc chấm dứt hợp đồng thực hiện theo căn cứ pháp luật và thỏa thuận.",
        "Thời hạn báo trước được áp dụng theo từng trường hợp.",
        "Người lao động hoàn tất bàn giao kho trước ngày cuối cùng.",
    )

    save(doc, "H05_Warehouse_Supervisor.docx")


# ============================================================
# H06
# ============================================================

def h06() -> None:
    doc = setup_document()
    title(doc, "H06 - Security Analyst")
    preamble(doc, "Security Analyst")

    heading(doc, 1, "Nhiệm vụ")
    paragraphs(
        doc,
        "Người lao động giám sát SOC, xử lý cảnh báo và phối hợp "
        "ứng phó sự cố an toàn thông tin.",
    )

    heading(doc, 2, "Thời gian đánh giá ban đầu")
    paragraphs(
        doc,
        "Trong 60 ngày đầu, người lao động thực hiện đầy đủ nhiệm vụ "
        "của vị trí Security Analyst.",
        "Kỹ năng phân tích sự cố và xử lý cảnh báo được đánh giá cuối giai đoạn.",
        "Nếu đáp ứng yêu cầu, người lao động tiếp tục đảm nhiệm vị trí.",
    )

    heading(doc, 3, "Thu nhập")
    paragraphs(
        doc,
        "Khoản thu nhập cố định là 24.000.000 đồng/tháng.",
        "Khoản này được thanh toán qua tài khoản vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Lịch công tác")
    paragraphs(
        doc,
        "Người lao động làm việc 08 giờ/ngày, từ thứ Hai đến thứ Sáu.",
        "Thời gian làm việc thông thường từ 09:00 đến 18:00.",
    )

    heading(doc, 5, "Quản lý tài khoản và thiết bị")
    paragraphs(
        doc,
        "Tài khoản đặc quyền và thiết bị được thu hồi "
        "khi quan hệ lao động kết thúc.",
    )

    heading(doc, 6, "Điều khoản áp dụng")
    paragraphs(
        doc,
        "Việc kết thúc hợp đồng được thực hiện theo Điều 8 của hợp đồng này.",
    )

    heading(doc, 7, "Bảo mật")
    paragraphs(
        doc,
        "Người lao động không được tiết lộ dữ liệu bảo mật "
        "hoặc thông tin cấu hình hệ thống.",
    )

    heading(doc, 8, "Quyền kết thúc quan hệ lao động")
    paragraphs(
        doc,
        "Mỗi bên có thể thực hiện quyền kết thúc quan hệ lao động "
        "khi có căn cứ phù hợp.",
        "Thời gian thông báo phải tuân thủ quy định áp dụng cho từng trường hợp.",
        "Người lao động phải hoàn tất bàn giao quyền truy cập và hồ sơ sự cố.",
    )

    save(doc, "H06_Security_Analyst.docx")


# ============================================================
# H07
# ============================================================

def h07() -> None:
    doc = setup_document()
    title(doc, "H07 - HR Specialist")
    preamble(doc, "HR Specialist")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động thực hiện tuyển dụng, hồ sơ nhân sự "
        "và hỗ trợ hoạt động quan hệ lao động.",
    )

    heading(doc, 2, "Thử việc")
    paragraphs(
        doc,
        "Thời gian thử việc là 60 ngày.",
        "Công việc trong thời gian thử việc tương ứng vị trí HR Specialist.",
        "Mức lương thử việc bằng 90% mức lương của vị trí.",
    )

    heading(doc, 3, "Đánh giá sau thử việc")
    paragraphs(
        doc,
        "Kết quả được đánh giá theo tiến độ tuyển dụng, độ chính xác hồ sơ "
        "và khả năng phối hợp.",
        "Công ty thông báo kết quả vào cuối giai đoạn.",
        "Nếu đạt yêu cầu, người lao động chuyển sang chế độ làm việc chính thức.",
    )

    heading(doc, 4, "Tiền lương")
    paragraphs(
        doc,
        "Mức lương theo công việc là 18.000.000 đồng/tháng.",
        "Tiền lương được thanh toán ngày 05 hằng tháng.",
    )

    heading(doc, 5, "Thời giờ làm việc")
    paragraphs(
        doc,
        "Thời giờ làm việc là 08 giờ/ngày, 05 ngày/tuần.",
        "Người lao động làm việc từ 08:30 đến 17:30.",
    )

    heading(doc, 6, "Nghỉ phép")
    paragraphs(
        doc,
        "Nghỉ phép năm được giải quyết theo chính sách và quy định pháp luật.",
    )

    heading(doc, 7, "Nghĩa vụ khi kết thúc")
    paragraphs(
        doc,
        "Khi chấm dứt, người lao động thực hiện việc bàn giao "
        "theo Điều 8 của hợp đồng này.",
    )

    heading(doc, 8, "Chấm dứt hợp đồng")
    paragraphs(
        doc,
        "Các trường hợp chấm dứt hợp đồng thực hiện theo quy định áp dụng.",
        "Thời hạn báo trước được xác định theo căn cứ và chủ thể thực hiện.",
    )

    save(doc, "H07_HR_Specialist.docx")


# ============================================================
# H08
# ============================================================

def h08() -> None:
    doc = setup_document()
    title(doc, "H08 - Sales Executive")
    preamble(doc, "Sales Executive")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động tìm kiếm khách hàng, tư vấn và thực hiện "
        "các hoạt động bán hàng.",
    )

    heading(doc, 2, "Giai đoạn thử việc")
    paragraphs(
        doc,
        "Thời gian thử việc là 45 ngày.",
        "Kết quả được đánh giá dựa trên khả năng tiếp cận và tư vấn khách hàng.",
    )

    heading(doc, 3, "Quyền lợi")
    table(
        doc,
        ["Nội dung", "Mức"],
        [
            ["Khoản cố định", "16.000.000 đồng/tháng"],
            ["Phụ cấp điện thoại", "500.000 đồng/tháng"],
            ["Hoa hồng", "Theo chính sách kinh doanh"],
        ],
    )
    paragraphs(
        doc,
        "Các khoản cố định được thanh toán vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Tổ chức công việc")
    paragraphs(
        doc,
        "Người lao động làm việc 08 giờ/ngày và 05 ngày/tuần.",
        "Lịch gặp khách hàng có thể được chủ động bố trí trong phạm vi công việc.",
    )

    heading(doc, 5, "Chỉ tiêu kinh doanh")
    paragraphs(
        doc,
        "Kết quả chỉ tiêu không làm thay đổi mức lương cố định "
        "nếu không có thỏa thuận khác bằng văn bản.",
    )

    heading(doc, 6, "Tài sản")
    paragraphs(
        doc,
        "Các thiết bị được giao phải được sử dụng đúng mục đích.",
    )

    heading(doc, 7, "Kết thúc hợp đồng")
    paragraphs(
        doc,
        "Hợp đồng có thể kết thúc theo các căn cứ và trình tự phù hợp.",
        "Bên thực hiện quyền kết thúc phải tuân thủ thời hạn thông báo.",
    )

    save(doc, "H08_Sales_Executive.docx")


# ============================================================
# H09
# ============================================================

def h09() -> None:
    doc = setup_document()
    title(doc, "H09 - DevOps Engineer")
    preamble(doc, "DevOps Engineer")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động vận hành CI/CD, hạ tầng cloud và hệ thống giám sát.",
    )

    heading(doc, 2, "Giai đoạn đầu")
    paragraphs(
        doc,
        "Trong hai tháng đầu kể từ ngày nhận việc, người lao động thực hiện "
        "nhiệm vụ triển khai và vận hành dưới sự đánh giá của quản lý kỹ thuật.",
        "Kết quả giai đoạn này được dùng để xác định việc tiếp tục bố trí "
        "người lao động vào vị trí.",
    )

    heading(doc, 3, "Chế độ")
    paragraphs(
        doc,
        "Khoản trả cố định cho công việc là 27.000.000 đồng mỗi tháng.",
        "Công ty chuyển khoản khoản này vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Lịch làm việc")
    paragraphs(
        doc,
        "Lịch làm việc thông thường từ 09:00 đến 18:00, 05 ngày/tuần.",
    )

    heading(doc, 5, "Trực ngoài giờ")
    paragraphs(
        doc,
        "Người lao động tham gia lịch on-call luân phiên theo tuần.",
        "Sau ca trực đêm, người lao động được bố trí khoảng nghỉ phù hợp "
        "trước khi trở lại lịch làm việc tiếp theo.",
    )

    heading(doc, 6, "Bảo mật và quyền truy cập")
    paragraphs(
        doc,
        "Thông tin tài khoản, khóa truy cập và bí mật hạ tầng phải được bảo mật.",
    )

    heading(doc, 7, "Điều khoản khác")
    paragraphs(
        doc,
        "Một bên có thể kết thúc quan hệ lao động trước thời hạn "
        "khi có căn cứ phù hợp.",
        "Thời gian thông báo và nghĩa vụ bàn giao được thực hiện "
        "theo quy định áp dụng.",
    )

    heading(doc, 8, "Bảo hiểm")
    paragraphs(
        doc,
        "Các chế độ bảo hiểm được thực hiện theo quy định.",
    )

    save(doc, "H09_DevOps_Engineer.docx")


# ============================================================
# H10 - PROBATION MISSING
# ============================================================

def h10() -> None:
    doc = setup_document()
    title(doc, "H10 - Office Administrator")
    preamble(doc, "Office Administrator")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động quản lý văn thư, lịch họp, hồ sơ hành chính "
        "và hỗ trợ vận hành văn phòng.",
    )

    heading(doc, 2, "Thời hạn hợp đồng")
    paragraphs(
        doc,
        "Hợp đồng xác định thời hạn 24 tháng kể từ ngày 01/11/2026.",
    )

    heading(doc, 3, "Tiền lương")
    paragraphs(
        doc,
        "Mức lương theo công việc là 14.500.000 đồng/tháng.",
        "Tiền lương được trả qua tài khoản vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Thời giờ làm việc")
    paragraphs(
        doc,
        "Người lao động làm việc 08 giờ/ngày, 05 ngày/tuần.",
        "Thời gian làm việc từ 08:00 đến 17:00.",
    )

    heading(doc, 5, "Tài sản")
    paragraphs(
        doc,
        "Tài sản và hồ sơ được giao phải hoàn trả khi hợp đồng chấm dứt.",
    )

    heading(doc, 6, "Điều khoản áp dụng")
    paragraphs(
        doc,
        "Các vấn đề về chấm dứt hợp đồng thực hiện theo Điều 8.",
    )

    heading(doc, 7, "Bảo hiểm và nghỉ phép")
    paragraphs(
        doc,
        "Người lao động tham gia các chế độ bảo hiểm và nghỉ phép năm "
        "theo quy định.",
    )

    heading(doc, 8, "Chấm dứt hợp đồng")
    paragraphs(
        doc,
        "Các trường hợp chấm dứt hợp đồng thực hiện theo quy định pháp luật.",
        "Thời hạn báo trước được áp dụng theo từng trường hợp.",
    )

    save(doc, "H10_Office_Administrator.docx")


# ============================================================
# H11
# ============================================================

def h11() -> None:
    doc = setup_document()
    title(doc, "H11 - QA Engineer")
    preamble(doc, "QA Engineer")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động thiết kế test case, kiểm thử sản phẩm "
        "và theo dõi lỗi phần mềm.",
    )

    heading(doc, 2, "Giai đoạn đánh giá ban đầu")
    paragraphs(
        doc,
        "Trong hai tháng đầu, người lao động thực hiện công việc QA "
        "trên dự án thực tế.",
        "Chất lượng test case, khả năng phát hiện lỗi và phối hợp "
        "được review cuối giai đoạn.",
        "Kết quả được dùng để quyết định việc tiếp tục bố trí vào vị trí.",
    )

    heading(doc, 3, "Tiền lương")
    table(
        doc,
        ["Khoản", "Giá trị"],
        [
            ["Thu nhập cố định", "18.500.000 đồng/tháng"],
            ["Phụ cấp", "600.000 đồng/tháng"],
        ],
    )
    paragraphs(
        doc,
        "Tiền lương được thanh toán vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Lịch làm việc")
    paragraphs(
        doc,
        "Người lao động làm việc từ 08:30 đến 17:30, "
        "từ thứ Hai đến thứ Sáu.",
    )

    heading(doc, 5, "Làm thêm và nghỉ bù")
    paragraphs(
        doc,
        "Khi phát sinh làm thêm theo kế hoạch phát hành, việc bố trí "
        "thời gian nghỉ bù được thực hiện theo kế hoạch của bộ phận.",
    )

    heading(doc, 6, "Thiết bị")
    paragraphs(
        doc,
        "Thiết bị kiểm thử phải được bảo quản và sử dụng đúng mục đích.",
    )

    heading(doc, 7, "Chấm dứt hợp đồng")
    paragraphs(
        doc,
        "Việc chấm dứt hợp đồng được thực hiện theo căn cứ "
        "và thời hạn báo trước tương ứng.",
    )

    save(doc, "H11_QA_Engineer.docx")


# ============================================================
# H12 - TERMINATION MISSING
# ============================================================

def h12() -> None:
    doc = setup_document()
    title(doc, "H12 - Finance Assistant")
    preamble(doc, "Finance Assistant")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động hỗ trợ nhập liệu kế toán, đối chiếu chứng từ "
        "và lập báo cáo nội bộ.",
    )

    heading(doc, 2, "Thử việc")
    paragraphs(
        doc,
        "Thời gian thử việc là 45 ngày.",
        "Kết quả thử việc được đánh giá theo độ chính xác "
        "và khả năng xử lý chứng từ.",
    )

    heading(doc, 3, "Chế độ")
    paragraphs(
        doc,
        "Mức lương theo công việc là 15.000.000 đồng/tháng.",
        "Phụ cấp ăn trưa là 500.000 đồng/tháng.",
        "Tiền lương được thanh toán vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Thời giờ làm việc")
    paragraphs(
        doc,
        "Thời giờ làm việc là 08 giờ/ngày và 05 ngày/tuần.",
        "Lịch làm việc từ 08:30 đến 17:30.",
    )

    heading(doc, 5, "Đánh giá hiệu quả")
    paragraphs(
        doc,
        "KPI được đánh giá định kỳ hằng quý sau khi người lao động "
        "đã làm việc chính thức.",
        "Nội dung đánh giá gồm độ chính xác, thời gian xử lý "
        "và mức độ tuân thủ quy trình.",
    )

    heading(doc, 6, "Bảo hiểm")
    paragraphs(
        doc,
        "Các chế độ bảo hiểm được thực hiện theo quy định.",
    )

    heading(doc, 7, "Bảo mật")
    paragraphs(
        doc,
        "Người lao động phải bảo mật dữ liệu tài chính "
        "và thông tin đối tác.",
    )

    heading(doc, 8, "Điều khoản chung")
    paragraphs(
        doc,
        "Hai bên có trách nhiệm thực hiện đúng các nội dung đã thống nhất.",
        "Mọi sửa đổi phải được lập thành văn bản.",
    )

    save(doc, "H12_Finance_Assistant.docx")


# ============================================================
# H13
# ============================================================

def h13() -> None:
    doc = setup_document()
    title(doc, "H13 - Call Center Agent")
    preamble(doc, "Call Center Agent")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động tiếp nhận và xử lý cuộc gọi của khách hàng.",
    )

    heading(doc, 2, "Giai đoạn đầu")
    paragraphs(
        doc,
        "Trong 30 ngày đầu, người lao động được đào tạo hệ thống "
        "và thực hiện cuộc gọi dưới sự giám sát.",
        "Kết quả về chất lượng cuộc gọi được dùng để xác định "
        "việc tiếp tục bố trí vào vị trí.",
    )

    heading(doc, 3, "Thu nhập")
    paragraphs(
        doc,
        "Khoản thu nhập cố định là 12.500.000 đồng/tháng.",
        "Khoản này được chuyển vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Phân ca")
    table(
        doc,
        ["Ca", "Giờ"],
        [
            ["A", "06:00-14:00"],
            ["B", "14:00-22:00"],
            ["C", "22:00-06:00"],
        ],
    )
    paragraphs(
        doc,
        "Lịch phân ca được công bố theo tuần.",
    )

    heading(doc, 5, "Nghỉ giữa ca và nghỉ hằng tuần")
    paragraphs(
        doc,
        "Người lao động được nghỉ giữa ca theo lịch vận hành.",
        "Ngày nghỉ hằng tuần được bố trí trong lịch phân ca.",
    )

    heading(doc, 6, "Tài sản")
    paragraphs(
        doc,
        "Tai nghe và thiết bị làm việc được quản lý theo quy định nội bộ.",
    )

    heading(doc, 7, "Điều khoản liên quan")
    paragraphs(
        doc,
        "Quyền kết thúc hợp đồng được thực hiện theo Điều 8.",
    )

    heading(doc, 8, "Kết thúc quan hệ lao động")
    paragraphs(
        doc,
        "Một bên có thể kết thúc quan hệ lao động khi có căn cứ phù hợp.",
        "Thời gian thông báo phải được tuân thủ theo từng trường hợp.",
    )

    save(doc, "H13_Call_Center_Agent.docx")


# ============================================================
# H14
# ============================================================

def h14() -> None:
    doc = setup_document()
    title(doc, "H14 - Procurement Specialist")
    preamble(doc, "Procurement Specialist")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động tìm kiếm nhà cung cấp, lấy báo giá "
        "và quản lý đơn mua hàng.",
    )

    heading(doc, 2, "Thử việc")
    paragraphs(
        doc,
        "Thời gian thử việc là 60 ngày.",
        "Kết quả thử việc căn cứ vào chất lượng hồ sơ mua hàng "
        "và khả năng phối hợp nhà cung cấp.",
    )

    heading(doc, 3, "Quyền lợi người lao động")
    paragraphs(
        doc,
        "Mức lương theo công việc là 19.000.000 đồng/tháng.",
        "Phụ cấp điện thoại là 400.000 đồng/tháng.",
        "Tiền lương được trả vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Tổ chức làm việc")
    paragraphs(
        doc,
        "Người lao động làm việc 08 giờ/ngày và 05 ngày/tuần.",
        "Thời gian làm việc từ 08:30 đến 17:30.",
    )

    heading(doc, 5, "Quy trình mua hàng")
    paragraphs(
        doc,
        "Việc thanh toán cho nhà cung cấp được thực hiện "
        "theo điều kiện trên đơn mua hàng.",
        "Bộ phận mua hàng có trách nhiệm theo dõi chứng từ thanh toán "
        "và tiến độ thanh toán cho nhà cung cấp.",
    )

    heading(doc, 6, "Bảo mật")
    paragraphs(
        doc,
        "Báo giá và thông tin thương mại của nhà cung cấp phải được bảo mật.",
    )

    heading(doc, 7, "Chấm dứt hợp đồng")
    paragraphs(
        doc,
        "Việc chấm dứt hợp đồng được thực hiện theo căn cứ "
        "và thời hạn báo trước áp dụng.",
    )

    save(doc, "H14_Procurement_Specialist.docx")


# ============================================================
# H15 - SALARY MISSING
# ============================================================

def h15() -> None:
    doc = setup_document()
    title(doc, "H15 - Data Engineer")
    preamble(doc, "Data Engineer")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động xây dựng pipeline dữ liệu, quản lý kho dữ liệu "
        "và hỗ trợ các hệ thống phân tích.",
    )

    heading(doc, 2, "Giai đoạn đánh giá ban đầu")
    paragraphs(
        doc,
        "Trong 60 ngày đầu, người lao động thực hiện nhiệm vụ Data Engineer "
        "trên hệ thống thực tế.",
        "Kết quả kỹ thuật và khả năng phối hợp được đánh giá cuối giai đoạn.",
        "Nếu đáp ứng yêu cầu, người lao động tiếp tục đảm nhiệm vị trí.",
    )

    heading(doc, 3, "Thời hạn hợp đồng")
    paragraphs(
        doc,
        "Hợp đồng có thời hạn 24 tháng kể từ ngày bắt đầu.",
    )

    heading(doc, 4, "Lịch công tác")
    paragraphs(
        doc,
        "Người lao động làm việc từ 09:00 đến 18:00, "
        "từ thứ Hai đến thứ Sáu.",
    )

    heading(doc, 5, "Trực hệ thống và nghỉ bù")
    paragraphs(
        doc,
        "Người lao động tham gia lịch trực hệ thống luân phiên.",
        "Sau thời gian trực ngoài giờ, bộ phận bố trí thời gian nghỉ bù phù hợp.",
    )

    heading(doc, 6, "Thiết bị và dữ liệu")
    paragraphs(
        doc,
        "Thiết bị phải hoàn trả khi quan hệ lao động kết thúc.",
        "Dữ liệu và khóa truy cập phải được bàn giao theo quy trình.",
    )

    heading(doc, 7, "Điều khoản áp dụng")
    paragraphs(
        doc,
        "Việc kết thúc quan hệ lao động được thực hiện theo Điều 8.",
    )

    heading(doc, 8, "Kết thúc quan hệ lao động")
    paragraphs(
        doc,
        "Một bên có thể kết thúc quan hệ lao động khi có căn cứ phù hợp.",
        "Thời gian thông báo được áp dụng theo từng trường hợp.",
    )

    save(doc, "H15_Data_Engineer.docx")


# ============================================================
# H16 - WORKING TIME MISSING
# ============================================================

def h16() -> None:
    doc = setup_document()
    title(doc, "H16 - Operations Coordinator")
    preamble(doc, "Operations Coordinator")

    heading(doc, 1, "Công việc")
    paragraphs(
        doc,
        "Người lao động điều phối hoạt động vận hành, xử lý yêu cầu nội bộ "
        "và theo dõi tiến độ công việc.",
    )

    heading(doc, 2, "Thử việc")
    paragraphs(
        doc,
        "Người lao động thử việc trong 45 ngày.",
        "Kết quả thử việc được đánh giá theo khả năng điều phối "
        "và độ chính xác báo cáo.",
    )

    heading(doc, 3, "Quyền lợi")
    table(
        doc,
        ["Khoản", "Mức"],
        [
            ["Lương theo công việc", "16.500.000 đồng/tháng"],
            ["Phụ cấp trách nhiệm", "1.000.000 đồng/tháng"],
            ["Hỗ trợ điện thoại", "400.000 đồng/tháng"],
        ],
    )
    paragraphs(
        doc,
        "Các khoản được thanh toán vào ngày 05 hằng tháng.",
    )

    heading(doc, 4, "Thời hạn hợp đồng")
    paragraphs(
        doc,
        "Hợp đồng xác định thời hạn 24 tháng kể từ ngày bắt đầu.",
    )

    heading(doc, 5, "Bảo hiểm và nghỉ phép năm")
    paragraphs(
        doc,
        "Người lao động tham gia các chế độ bảo hiểm theo quy định.",
        "Nghỉ phép năm được giải quyết theo chính sách của công ty.",
    )

    heading(doc, 6, "Tài sản")
    paragraphs(
        doc,
        "Tài sản được giao phải được bảo quản theo quy định nội bộ.",
    )

    heading(doc, 7, "Chấm dứt hợp đồng")
    paragraphs(
        doc,
        "Việc chấm dứt hợp đồng thực hiện theo căn cứ pháp luật "
        "và thỏa thuận của hai bên.",
        "Thời hạn báo trước được áp dụng theo từng trường hợp.",
    )

    heading(doc, 8, "Điều khoản chung")
    paragraphs(
        doc,
        "Các sửa đổi hoặc bổ sung hợp đồng phải được lập thành văn bản.",
    )

    save(doc, "H16_Operations_Coordinator.docx")


def main() -> None:
    generators = [
        h01,
        h02,
        h03,
        h04,
        h05,
        h06,
        h07,
        h08,
        h09,
        h10,
        h11,
        h12,
        h13,
        h14,
        h15,
        h16,
    ]

    for generator in generators:
        generator()

    files = sorted(ROOT.glob("*.docx"))

    if len(files) != 16:
        raise RuntimeError(
            f"Expected 16 DOCX files, got {len(files)}"
        )

    print()
    print(f"Generated {len(files)} held-out contracts.")


if __name__ == "__main__":
    main()
